"""Tests for Word ``.docx`` ingestion (RAG R4).

docx text extraction is monkeypatched for deterministic unit tests (the real
backend — python-docx — is opt-in and exercised by the skipif-guarded live
test), mirroring how PDF/OCR are stubbed in ``test_pdf_ocr.py``.
"""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

from iris_harness.services.rag import ingest as ingest_mod
from iris_harness.services.rag.ingest import ingest_path
from iris_harness.services.rag.qa import answer_question
from iris_harness.services.rag.retrieve import search_documents
from iris_harness.services.rag.store import DocumentStore

_DOCX_AVAILABLE = importlib.util.find_spec("docx") is not None


@pytest.fixture
def store(tmp_path: Path) -> DocumentStore:
    s = DocumentStore(db_path=tmp_path / "rag.db")
    s.ensure_schema()
    return s


# ─── docx (text extraction monkeypatched) ───────────────────────────────


def test_docx_ingested(
    store: DocumentStore, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    doc = tmp_path / "proposal.docx"
    doc.write_bytes(b"PK fake docx bytes")
    monkeypatch.setattr(
        ingest_mod,
        "extract_docx_text",
        lambda path: "Project goals for Q2.\nWe will ship the docx loader.",
    )
    r = ingest_path(doc, store=store, index=None)
    assert r.sources_added == 1 and r.chunks_indexed == 1
    src = store.list_sources()[0]
    assert src.kind == "file" and src.title == "proposal"

    ans = answer_question("what will we ship?", store=store, index=None)
    assert "docx loader" in ans.answer


def test_docx_title_is_file_stem(
    store: DocumentStore, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    doc = tmp_path / "Q2 Plan.docx"
    doc.write_bytes(b"PK fake")
    monkeypatch.setattr(ingest_mod, "extract_docx_text", lambda path: "Some content here.")
    ingest_path(doc, store=store, index=None)
    hits = search_documents("content", store=store, index=None)
    assert hits and "Q2 Plan.docx#0" in hits[0].citation.label()


def test_empty_docx_skipped(
    store: DocumentStore, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    doc = tmp_path / "blank.docx"
    doc.write_bytes(b"PK fake")
    monkeypatch.setattr(ingest_mod, "extract_docx_text", lambda path: "   \n  ")
    r = ingest_path(doc, store=store, index=None)
    assert r.sources_added == 0 and r.sources_skipped == 1  # nothing indexable


def test_docx_skipped_without_python_docx(
    store: DocumentStore, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    doc = tmp_path / "report.docx"
    doc.write_bytes(b"PK fake")

    def _missing(path: object) -> str:
        raise RuntimeError("install python-docx to ingest .docx")

    monkeypatch.setattr(ingest_mod, "extract_docx_text", _missing)
    r = ingest_path(doc, store=store, index=None)
    assert r.sources_added == 0 and r.sources_skipped == 1  # graceful, no crash


# ─── real python-docx (only where it is installed) ───────────────────────


@pytest.mark.skipif(not _DOCX_AVAILABLE, reason="python-docx not installed")
def test_real_docx_roundtrip(store: DocumentStore, tmp_path: Path) -> None:
    import docx  # type: ignore[import-untyped]

    doc_path = tmp_path / "notes.docx"
    document = docx.Document()
    document.add_paragraph("Invoice total is 4242 USD.")
    document.add_paragraph("Payment is due next month.")
    document.save(str(doc_path))

    r = ingest_path(doc_path, store=store, index=None)
    assert r.sources_added == 1
    hits = search_documents("invoice total", store=store, index=None)
    assert hits and "4242" in hits[0].text
