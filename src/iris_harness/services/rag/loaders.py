"""Document text loaders — PDF (RAG R2) and Word ``.docx`` (RAG R4).

PDFs return text **per page** so chunks can cite the page they came from, via
``pypdfium2`` (the permissively-licensed backend IRIS already bundles for
Finance) — *not* PyMuPDF/pymupdf4llm, which are AGPL and deliberately not
bundled in an Apache-2.0 project. ``.docx`` returns the full document text via
``python-docx``, lazy-imported and opt-in (mirroring OCR). Local-only; the
file never leaves the box.
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_pdf_pages(path: str | Path, *, password: str | None = None) -> list[str]:
    """Extract text from a PDF, one string per page (page 1 → index 0)."""
    import pypdfium2 as pdfium  # type: ignore[import-untyped]

    pdf = pdfium.PdfDocument(str(path), password=password or None)
    try:
        pages: list[str] = []
        for page in pdf:
            textpage = page.get_textpage()
            try:
                pages.append(textpage.get_text_range() or "")
            finally:
                textpage.close()
                page.close()
        return pages
    finally:
        pdf.close()


def extract_docx_text(path: str | Path) -> str:
    """Extract text from a Microsoft Word ``.docx``, paragraphs joined by ``\\n``.

    ``python-docx`` is **opt-in** (not bundled, like OCR's pytesseract): it is
    lazy-imported here so the base install stays light, and a missing install
    raises a clear ``RuntimeError`` that the ingest layer turns into a logged
    skip rather than a crash. Local-only; the file never leaves the box.
    """
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError(
            "install python-docx to ingest .docx (`pip install python-docx`)"
        ) from exc

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def render_pdf_page(
    path: str | Path, page_index: int, *, password: str | None = None, scale: float = 2.0
) -> object:
    """Render one PDF page to a PIL image — for OCR of scanned/image-only pages.

    ``scale`` 2.0 ≈ 144 DPI, a good speed/accuracy trade-off for OCR.
    """
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path), password=password or None)
    try:
        page = pdf[page_index]
        try:
            return page.render(scale=scale).to_pil()
        finally:
            page.close()
    finally:
        pdf.close()


__all__ = ["extract_docx_text", "extract_pdf_pages", "render_pdf_page"]
